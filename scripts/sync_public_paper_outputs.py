#!/usr/bin/env python3
"""Extract the public-safe paper audit layer from the frozen reader bundle.

This command is intentionally one way: it reads a manuscript reader bundle and
its Supplementary Data workbook, then writes only an explicit allow-list of
aggregate sheets and non-station-resolved reference figures.  Raw observations,
station identifiers, coordinates, station-time series and stationwise maps are
never copied.

The command is useful to the authors when a manuscript result is frozen.  A
public user normally consumes the already bundled ``data/paper_aggregates``
tables and renders them with ``r2p-paper-figures``.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from scrub_office_metadata import scrub_office_file


SAFE_SHEETS = (
    "Table1_main",
    "TableS1a_training",
    "TableS1b_readout",
    "TableS2_cases",
    "Case_definitions",
    "Table1_seed_numeric",
    "Gauge_contract_matrix",
    "Fig2a_density",
    "Fig2a_summary",
    "Fig2b_amount_bins",
    "Fig2c_fixed_metrics",
    "Fig2d_CSI_summary",
    "Fig2d_FB_summary",
    "Fig2d_readout_members",
    "Fig2d_delta_CI",
    "Categorical_per_rep",
    "Categorical_mean_sd",
    "Continuous_per_rep",
    "Continuous_mean_sd",
    "Paired_date_bootstrap",
    "Fig3ef_route_contrasts",
    "FigS1_distance_cells",
    "FigS1_distance_summary",
    "Allstation_overall",
    "Allstation_seedwise",
    "Route_metric_CI",
    "Fig5_intensity_point_CI",
    "TableS2b_diagnostics",
    "Case_per_rep",
    "Case_mean_sd",
    "FigS3_pairwise_counts",
    "Field_instant_CSIM",
    "Field_instant_CSI",
    "Field_RN60_CSIM",
    "Field_RN60_CSI",
    "Field_RN60_CSI_delta_CI",
)

# Treat the export as a schema-bound data product, not merely a list of sheet
# names.  This prevents an upstream workbook from silently adding a station
# identifier, coordinate, timestamp or other newly sensitive column to a
# previously approved sheet.
SAFE_SHEET_HEADERS = {
    "Allstation_overall": ("contrast", "left_route", "right_route", "lead_min", "threshold_mm", "allstation_csi", "direct_r2p_csi", "delta_csi", "ci_low", "ci_high", "n_finite_common", "bootstrap_reps", "member_count", "member_ids"),
    "Allstation_seedwise": ("contrast", "seed", "lead_min", "threshold_mm", "allstation_csi", "direct_r2p_csi", "delta_csi", "allstation_hit", "allstation_miss", "allstation_false_alarm", "direct_r2p_hit", "direct_r2p_miss", "direct_r2p_false_alarm", "n_finite_common"),
    "Case_definitions": ("case_rank", "window_start_exclusive", "window_end_inclusive", "q95_station_total_mm", "n_complete_stations", "selected", "mechanism_label", "event_label"),
    "Case_mean_sd": ("case_rank", "event_label", "lead_min", "threshold_mm", "route", "route_label", "replicate_count", "csi_mean", "csi_sd", "csi_min", "csi_max", "finite_pair_count", "n_issue_times"),
    "Case_per_rep": ("case_rank", "event_label", "lead_min", "threshold_mm", "route", "route_label", "replicate", "replicate_semantics", "finite_pair_count", "n_issue_times", "hit", "miss", "false_alarm", "correct_negative", "csi"),
    "Categorical_mean_sd": ("route_key", "route_label", "lead_min", "threshold_mm", "replicate_count", "hit_mean", "hit_sd", "miss_mean", "miss_sd", "false_alarm_mean", "false_alarm_sd", "correct_negative_mean", "correct_negative_sd", "csi_mean", "csi_sd", "pod_mean", "pod_sd", "far_mean", "far_sd", "frequency_bias_mean", "frequency_bias_sd", "finite_pair_count_mean", "finite_pair_count_sd"),
    "Categorical_per_rep": ("route_key", "route_label", "replicate", "replicate_id", "lead_min", "threshold_mm", "hit", "miss", "false_alarm", "correct_negative", "csi", "pod", "far", "frequency_bias", "finite_pair_count"),
    "Continuous_mean_sd": ("route_key", "route_label", "lead_min", "replicate_count", "rmse_mean", "rmse_sd", "mean_bias_mean", "mean_bias_sd", "correlation_mean", "correlation_sd"),
    "Continuous_per_rep": ("route_key", "route_label", "replicate", "replicate_id", "lead_min", "rmse", "mean_bias", "correlation", "n", "sum_prediction", "sum_truth", "sum_prediction_sq", "sum_truth_sq", "sum_product", "sum_error", "sum_squared_error"),
    "Field_RN60_CSI": ("model", "lead_min", "threshold_mm", "csi"),
    "Field_RN60_CSIM": ("model", "lead_min", "csi_m"),
    "Field_RN60_CSI_delta_CI": ("lead_min", "threshold_mm", "long_csi", "pysteps_csi", "delta_long_minus_pysteps", "bootstrap_mean_delta", "ci95_lower", "ci95_upper", "probability_delta_gt_zero", "ci_excludes_zero"),
    "Field_instant_CSI": ("model", "lead_min", "threshold_mm_h", "csi"),
    "Field_instant_CSIM": ("model", "lead_min", "csi_m"),
    "Fig2a_density": ("gauge_log1p", "fixed_readout_log1p", "pair_count"),
    "Fig2a_summary": ("support_definition", "pearson_r", "support_pairs", "finite_pairs", "density_display_pairs"),
    "Fig2b_amount_bins": ("gauge_bin", "n_pairs", "fixed_readout_to_gauge_amount_ratio", "gauge_conditioned_amount_bias_mm"),
    "Fig2c_fixed_metrics": ("threshold_mm", "hits", "misses", "false_alarms", "csi", "frequency_bias", "n_finite"),
    "Fig2d_CSI_summary": ("route", "threshold_mm", "csi_mean", "csi_sd"),
    "Fig2d_FB_summary": ("route", "threshold_mm", "frequency_bias_mean", "frequency_bias_sd", "n_members"),
    "Fig2d_delta_CI": ("family", "threshold_mm", "delta_csi_mean", "ci_low", "ci_high", "bootstrap_reps", "block", "n_dates", "interval"),
    "Fig2d_readout_members": ("route", "replicate", "threshold_mm", "hits", "misses", "false_alarms", "csi", "frequency_bias", "finite_pairs"),
    "Fig3ef_route_contrasts": ("panel", "contrast", "left_route", "right_route", "lead_min", "threshold_mm", "left_csi", "right_csi", "delta_csi", "ci_low", "ci_high", "n_finite_common", "bootstrap_reps"),
    "Fig5_intensity_point_CI": ("route", "route_label", "replicate", "aggregation", "bootstrap_reps", "bootstrap_unit", "ci_definition", "lead_min", "metric", "x_index", "x_label", "bin_low_mm", "bin_high_mm", "threshold_mm", "estimate", "ci_low", "ci_high", "n_common_pairs", "n_observed_events"),
    "FigS1_distance_cells": ("Distance group", "Distance to nearest fitting station (km)", "Stations", "Lead (min)", "RN60 threshold (mm)", "All-station R2P CSI", "Direct R2P CSI", "ΔCSI (All-station − Direct)", "95% CI low", "95% CI high", "Common finite pairs"),
    "FigS1_distance_summary": ("distance_group", "summary", "threshold_mm", "lead_min", "delta_csi", "ci_low", "ci_high", "averaged_lead_count", "averaged_threshold_count"),
    "FigS3_pairwise_counts": ("threshold_mm", "comparison", "stations_left_higher", "stations_equal", "stations_left_lower", "stations_total", "median_stationwise_csi_difference", "mean_stationwise_csi_difference"),
    "Gauge_contract_matrix": ("Evaluation setting", "Route / experiment", "Fitting set (n=514)—targets for fitting or selection", "Held-out set (n=128)—targets for fitting or selection", "Fitting set (n=514)—inputs at issuance", "Held-out set (n=128)—inputs at issuance"),
    "Paired_date_bootstrap": ("metric_group", "metric", "contrast", "left_route", "right_route", "lead_min", "threshold_mm", "left_point", "right_point", "delta", "ci_low", "ci_high", "ci_excludes_zero", "bootstrap_valid_reps", "bootstrap_reps", "member_count", "member_ids"),
    "Route_metric_CI": ("metric_group", "metric", "route_key", "route_label", "lead_min", "threshold_mm", "point_estimate", "ci_low", "ci_high", "bootstrap_valid_reps", "bootstrap_reps"),
    "Table1_main": ("Route", "Lead (min)", "RN60 CSI-M", "CSI ≥1 mm", "CSI ≥5 mm", "CSI ≥10 mm", "CSI ≥20 mm", "RMSE (mm)", "Mean bias (mm)", "Correlation"),
    "Table1_seed_numeric": ("route_key", "route_label", "replicate", "replicate_id", "lead_min", "csi_m", "rmse", "mean_bias", "correlation", "csi_1mm", "pod_1mm", "far_1mm", "frequency_bias_1mm", "hit_1mm", "miss_1mm", "false_alarm_1mm", "correct_negative_1mm", "csi_5mm", "pod_5mm", "far_5mm", "frequency_bias_5mm", "hit_5mm", "miss_5mm", "false_alarm_5mm", "correct_negative_5mm", "csi_10mm", "pod_10mm", "far_10mm", "frequency_bias_10mm", "hit_10mm", "miss_10mm", "false_alarm_10mm", "correct_negative_10mm", "csi_20mm", "pod_20mm", "far_20mm", "frequency_bias_20mm", "hit_20mm", "miss_20mm", "false_alarm_20mm", "correct_negative_20mm"),
    "TableS1a_training": ("Route / control", "Fitting data", "Selection", "Final members", "Objective / defining change"),
    "TableS1b_readout": ("Field source", "3 × 3 Patch MLP (macro CSI; epoch)", "3 × 3 Patch CNN (macro CSI; epoch)"),
    "TableS2_cases": ("Event window (KST)", "Across-station 95th-percentile total (mm)", "Complete stations (n)"),
    "TableS2b_diagnostics": ("Route", "Lead (min)", "RN60 threshold (mm)", "POD", "FAR", "Frequency bias", "Hit", "Miss", "False alarm", "Correct negative", "Finite pairs"),
}

# These manuscript assets contain no station identifiers, coordinates or
# station-time observations.  Fig. 1 and Supplementary Figs. S2/S3 are omitted
# pending the separate station-artifact redistribution decision.
SAFE_REFERENCE_ASSETS = {
    "02_Figure_2_valid_time_readout.png": "figure2_valid_time_readout.png",
    "03_Figure_3_route_skill_and_contrasts.png": "figure3_route_skill_and_contrasts.png",
    "04_Figure_4_R2P_ablations.png": "figure4_r2p_ablations.png",
    "05_Figure_5_heavy_rain_attenuation.png": "figure5_heavy_rain_attenuation.png",
    "06_Figure_6_route_architecture.png": "figure6_route_architecture.png",
    "07_Supplementary_Figure_S1_allstation_distance_diagnostic.png": "figureS1_allstation_distance_diagnostic.png",
    "10_Supplementary_Figure_S4_native_field_verification.png": "figureS4_native_field_verification.png",
    "11_Table_1_heldout_point_skill_EN.png": "table1_heldout_point_skill.png",
}

FORBIDDEN_COLUMNS = {
    "station_id",
    "station_name",
    "name",
    "lat",
    "latitude",
    "lon",
    "longitude",
    "valid_time_kst",
    "issue_time_kst",
    "station_code",
    "stn_id",
    "stnid",
    "aws_id",
    "aws_code",
    "site_id",
    "site_code",
    "지점",
    "지점명",
    "지점번호",
    "위도",
    "경도",
}


class PublicSyncError(RuntimeError):
    """Raised when the proposed public export departs from the allow-list."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(8 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_cell(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def rows_from_sheet(sheet: Any) -> tuple[list[str], list[list[Any]]]:
    iterator = sheet.iter_rows(values_only=True)
    try:
        raw_header = next(iterator)
    except StopIteration as exc:
        raise PublicSyncError(f"empty sheet: {sheet.title}") from exc
    header = [str(value).strip() if value is not None else "" for value in raw_header]
    if not header or any(not value for value in header):
        raise PublicSyncError(f"blank or missing header in {sheet.title}")
    expected_header = SAFE_SHEET_HEADERS.get(sheet.title)
    if expected_header is None or tuple(header) != expected_header:
        raise PublicSyncError(
            f"header drift in {sheet.title}: expected {expected_header!r}, found {tuple(header)!r}"
        )
    normalized = {value.strip().lower().replace(" ", "_") for value in header}
    blocked = sorted(normalized & FORBIDDEN_COLUMNS)
    if blocked:
        raise PublicSyncError(
            f"allow-listed sheet {sheet.title} unexpectedly contains restricted columns: {blocked}"
        )
    rows = [[normalize_cell(value) for value in row] for row in iterator]
    return header, rows


def write_csv(path: Path, header: list[str], rows: Iterable[list[Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".partial")
    with temporary.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)
    temporary.replace(path)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".partial")
    temporary.write_text(
        json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def export_workbook(workbook_path: Path, output_root: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:  # pragma: no cover - dependency message
        raise SystemExit(
            "openpyxl is required only for author-side synchronization; "
            "install the project paper-output optional dependencies"
        ) from exc

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    missing = sorted(set(SAFE_SHEETS) - set(workbook.sheetnames))
    if missing:
        raise PublicSyncError(f"Supplementary Data workbook lacks sheets: {missing}")

    records: list[dict[str, Any]] = []
    public_book = Workbook(write_only=True)
    # Make the convenience XLSX byte-stable.  openpyxl otherwise records the
    # wall-clock build time in the core properties and ZIP entry metadata.
    frozen_time = datetime(2000, 1, 1, 0, 0, 0)
    public_book.properties.created = frozen_time
    public_book.properties.modified = frozen_time
    for sheet_name in SAFE_SHEETS:
        header, rows = rows_from_sheet(workbook[sheet_name])
        csv_path = output_root / f"{sheet_name}.csv"
        write_csv(csv_path, header, rows)
        records.append(
            {
                "kind": "aggregate_sheet",
                "name": sheet_name,
                "path": csv_path.name,
                "rows": len(rows),
                "columns": len(header),
                "sha256": sha256_file(csv_path),
            }
        )

        target = public_book.create_sheet(title=sheet_name)
        target.append(header)
        for row in rows:
            target.append(row)

    aggregate_workbook = output_root / "Supplementary_Data_1_public_aggregate.xlsx"
    public_book.save(aggregate_workbook)
    scrub_office_file(aggregate_workbook)
    records.append(
        {
            "kind": "aggregate_workbook",
            "name": "Supplementary Data 1 public aggregate subset",
            "path": aggregate_workbook.name,
            "sha256": sha256_file(aggregate_workbook),
            "sheet_count": len(SAFE_SHEETS),
        }
    )
    return records


def copy_reference_assets(reader_assets: Path, output_dir: Path) -> list[dict[str, Any]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []
    for source_name, public_name in SAFE_REFERENCE_ASSETS.items():
        source = reader_assets / source_name
        if not source.is_file():
            raise FileNotFoundError(source)
        destination = output_dir / public_name
        shutil.copyfile(source, destination)
        records.append(
            {
                "kind": "reference_rendering",
                "name": public_name,
                "path": str(Path("figures/reference") / public_name),
                "sha256": sha256_file(destination),
            }
        )
    return records


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument("--supplementary-data", type=Path, required=True)
    result.add_argument("--reader-assets", type=Path, required=True)
    result.add_argument(
        "--release-root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
        help="release_4km10min root (defaults to this script's parent repository)",
    )
    return result


def main() -> None:
    args = parser().parse_args()
    release_root = args.release_root.expanduser().resolve()
    workbook = args.supplementary_data.expanduser().resolve()
    reader_assets = args.reader_assets.expanduser().resolve()
    if not workbook.is_file() or not reader_assets.is_dir():
        raise FileNotFoundError("both --supplementary-data and --reader-assets must exist")

    aggregate_root = release_root / "data" / "paper_aggregates"
    aggregate_root.mkdir(parents=True, exist_ok=True)
    records = export_workbook(workbook, aggregate_root)
    records += copy_reference_assets(reader_assets, release_root / "figures" / "reference")
    manifest = {
        "schema": "r2p-4km10min-public-paper-audit-layer-v1",
        "source_supplementary_data_sha256": sha256_file(workbook),
        "source_reader_manifest_sha256": sha256_file(reader_assets / "00_ASSET_MANIFEST.json"),
        "station_resolved_content_included": False,
        "excluded_reader_assets": [
            "Figure 1 (station coordinates and radar footprint)",
            "Supplementary Figure S2 (station time series)",
            "Supplementary Figure S3 (stationwise map)",
        ],
        "records": records,
    }
    write_json(aggregate_root / "manifest.json", manifest)
    print(f"wrote {len(records)} public-safe aggregate/reference records")


if __name__ == "__main__":
    main()
